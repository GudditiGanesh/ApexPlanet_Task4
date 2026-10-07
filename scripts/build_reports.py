"""Task 4 | Builds report/Task4_Report.pdf, report/Hypothesis_Test_Report.pdf, report/Business_Insights.docx."""
import json
from pathlib import Path

import matplotlib
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
R = json.loads((ROOT / "data" / "results.json").read_text())
K, H = R["kpis"], R["hypothesis"]
NAVY = colors.HexColor("#1F3A5F")
M = lambda v, d=1: f"₹{v / 1e6:.{d}f}M"
Kf = lambda v: ("−" if v < 0 else "") + f"₹{abs(v) / 1e3:.1f}K"
pct = lambda v, d=0: f"{v * 100:.{d}f}%"

fd = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
pdfmetrics.registerFont(TTFont("DV", str(fd / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("DV-B", str(fd / "DejaVuSans-Bold.ttf")))
pdfmetrics.registerFontFamily("DV", normal="DV", bold="DV-B", italic="DV", boldItalic="DV-B")
ss = getSampleStyleSheet()
BODY = ParagraphStyle("body", parent=ss["Normal"], fontName="DV", fontSize=9.5, leading=14, spaceAfter=6)
H1 = ParagraphStyle("h1", parent=BODY, fontName="DV-B", fontSize=15, textColor=NAVY, spaceBefore=10, spaceAfter=8)
TITLE = ParagraphStyle("t", parent=H1, fontSize=22, leading=28, spaceAfter=4)
SMALL = ParagraphStyle("s", parent=BODY, fontSize=8.5, leading=11, spaceAfter=0)

# ---- derived numbers ----
mm = [r for r in R["monthly"] if r["Complete_Month"]]
avg_m = sum(r["Revenue"] for r in mm) / len(mm)
aug, sep = (next(r["Revenue"] for r in mm if r["Month"] == f"2025-{n}") for n in ("08", "09"))
dip_gain = (avg_m - aug) + (avg_m - sep)
retention_gain = 0.05 * K["Unique Customers"] * K["Average Order Value"]
city = {c["City"]: c for c in R["city"]}
top3_orders = sum(c["Orders"] for c in R["city"][:3]) / 3
pune_gain = (top3_orders - city["Pune"]["Orders"]) * city["Pune"]["AOV"]
seg = {s["Segment"]: s for s in R["segments"]}
cat = {c["Category"]: c for c in R["category"]}


def img(name, w_cm=16):
    p = ROOT / "charts" / name
    w, h = PILImage.open(p).size
    return Image(str(p), width=w_cm * cm, height=w_cm * cm * h / w)


def tbl(rows, widths, header=True):
    rows = [[Paragraph(str(c), SMALL if i else ParagraphStyle("hd", parent=SMALL, fontName="DV-B", textColor=colors.white)) for c in r] for i, r in enumerate(rows)]
    t = Table(rows, colWidths=[w * cm for w in widths], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DBE3")),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F8")]),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    return t


def build(path, story):
    def foot(c, d):
        c.setFont("DV", 8); c.setFillColor(colors.grey)
        c.drawString(2 * cm, 1.2 * cm, "ApexPlanet Internship | Task 4: Data Storytelling & Statistical Validation"); c.drawRightString(19 * cm, 1.2 * cm, str(d.page))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=2 * cm,
                      title=path.stem.replace("_", " ")).build(story, onFirstPage=foot, onLaterPages=foot)


P = lambda t, s=BODY: Paragraph(t, s)

# =========================== Task4_Report.pdf ===========================
story = [P("From Data to Decisions: The 2025 Sales Story", TITLE), P("Business story report | ApexPlanet Internship, Task 4", BODY),
         P("1. Executive summary", H1),
         P(f"The business earned <b>{M(K['Total Revenue'])}</b> from <b>{K['Total Orders']:,}</b> orders placed by <b>{K['Unique Customers']}</b> customers in 2025. "
           f"Electronics is the largest category ({pct(K['Top Category Revenue Share'])} of revenue). A hypothesis test asked whether Electronics orders are significantly larger than Fashion orders: "
           f"the answer is <b>not at the 5% level</b> (p = {H['welch_p']:.3f}), so Electronics' lead comes from <b>order volume</b>, not larger orders. "
           f"The clearest opportunities are <b>customer retention</b> (only {pct(K['Repeat Customer Rate'])} of customers reorder) and the <b>August–September dip</b>."),
         P("2. Dataset and cleaning", H1),
         P(f"The dataset has 1,000 orders and 12 columns covering {R['quality']['date_range'][0]} to {R['quality']['date_range'][1]}. Quality checks and actions:"),
         tbl([["Issue", "Records", "Action"], ["Missing Age", R["quality"]["missing_age_imputed"], "Filled with median age"],
              ["Missing City", R["quality"]["missing_city_labelled_unknown"], "Labelled 'Unknown'; excluded from city ranking"],
              ["Order_ID ORD100050 reused for different orders", R["quality"]["reused_order_ids"], "Kept as distinct orders and flagged"],
              ["Fully duplicated rows", 0, "None found"], ["Total_Sales = Quantity × Unit_Price", "1,000 of 1,000", "Verified"],
              ["Partial month (1 Jan 2026)", "5 orders", "In totals; excluded from monthly trend"]], [7, 2.5, 7.5]),
         P("3. Key performance indicators", H1),
         tbl([["KPI", "Value"], ["Total revenue", M(K["Total Revenue"])], ["Orders / customers", f"{K['Total Orders']:,} / {K['Unique Customers']}"],
              ["Average / median order value", f"{Kf(K['Average Order Value'])} / {Kf(K['Median Order Value'])}"], ["Units sold", f"{K['Units Sold']:,}"],
              ["Top category / city / product", f"{K['Top Category']} / {K['Top City']} / {K['Top Product']}"], ["Repeat customer rate", pct(K["Repeat Customer Rate"])]], [7, 10]),
         Spacer(1, 8), img("dashboard_summary.png", 16), PageBreak(),
         P("4. Sales trend", H1),
         P(f"Monthly revenue averaged {M(avg_m)}, peaking in March ({M(max(r['Revenue'] for r in mm))}) and bottoming in September ({M(sep)}, {pct(1 - sep / avg_m)} below average). Revenue recovered from October."),
         img("sales_trend.png", 13),
         P("5. Customer analysis", H1),
         P(f"Customers were segmented by lifetime spend. The <b>High Value</b> segment ({seg['High Value']['Customers']} customers, 20%) generates {pct(seg['High Value']['Revenue_Share'])} of revenue, while the "
           f"<b>Low Value</b> segment (40% of customers) generates only {pct(seg['Low Value']['Revenue_Share'])}. Because {pct(1 - K['Repeat Customer Rate'])} of customers bought once, growth depends on retention."),
         img("customer_segments.png", 11), PageBreak(),
         P("6. Product and market performance", H1),
         P(f"Electronics earned {M(cat['Electronics']['Revenue'])} from {cat['Electronics']['Orders']} orders (two products, Laptop and Mobile). Every other category is a single product, so category revenue largely tracks order count; "
           f"average order value is similar everywhere (ANOVA p = {H['anova_p']:.2f}). Top cities Patna, Kolkata, Bengaluru and Mumbai are within ₹0.5M of each other; Gaya and Pune trail."),
         img("category_sales.png", 16), img("city_sales.png", 10),
         P("7. Statistical validation", H1),
         P(f"<b>H0:</b> mean order value (Electronics) = Fashion. <b>H1:</b> Electronics &gt; Fashion. Welch t = {H['welch_t']:.2f}, df = {H['welch_df']:.0f}, <b>p = {H['welch_p']:.3f}</b>; "
           f"difference {Kf(H['diff'])} ({pct(H['diff_pct'], 1)}), 95% CI {Kf(H['diff_ci95'][0])} to {Kf(H['diff_ci95'][1])}. <b>Decision: {H['decision']}</b> at α = 0.05. See the Hypothesis Test Report for details."),
         P("8. Recommendations and expected impact", H1),
         tbl([["Recommendation", "Evidence", "Expected impact (illustrative)"],
              ["Retain high-value customers: loyalty offers, reorder reminders", f"20% of customers give {pct(seg['High Value']['Revenue_Share'])} of revenue; {pct(K['Repeat Customer Rate'])} repeat rate", f"Repeat rate 5% to 10% adds about {M(retention_gain)} (+{pct(retention_gain / K['Total Revenue'], 1)})"],
              ["Run promotions in August and September", f"{M(aug)} and {M(sep)} vs {M(avg_m)} average", f"Reaching the average in both months adds about {M(dip_gain)}"],
              ["Back Electronics through volume, not price assumptions", "Lead is order count; order-size gap not significant", "Avoids unproven premium pricing; re-test with more data"],
              ["Grow Pune", f"High AOV ({Kf(city['Pune']['AOV'])}) but only {city['Pune']['Orders']} orders", f"Matching top-3 city volume adds about {M(pune_gain)}"]], [5.5, 5.5, 6]),
         P("9. Conclusion", H1),
         P("Revenue is healthy and Electronics leads, but the data does not prove Electronics orders are larger than Fashion orders. The evidence-backed levers are retention, the late-summer dip and under-penetrated cities. "
           "Impact figures are illustrative scenarios, not forecasts.")]
build(ROOT / "report" / "Task4_Report.pdf", story)

# ======================= Hypothesis_Test_Report.pdf =======================
fmt_ci = lambda ci: f"{Kf(ci[0])} to {Kf(ci[1])}"
story = [P("Hypothesis Test Report", TITLE), P("Do Electronics orders generate significantly higher sales than Fashion orders?", BODY),
         P("1. Business question and hypotheses", H1),
         P("<b>H0:</b> Average order value (Electronics) = Average order value (Fashion).<br/><b>H1:</b> Average order value (Electronics) &gt; Average order value (Fashion) (one-sided).<br/>Significance level α = 0.05. Unit of analysis: one order (Total_Sales)."),
         P("2. Data", H1),
         tbl([["", "Electronics", "Fashion"], ["Orders (n)", H["n_electronics"], H["n_fashion"]], ["Mean order value", Kf(H["mean_electronics"]), Kf(H["mean_fashion"])],
              ["Median order value", Kf(H["median_electronics"]), Kf(H["median_fashion"])], ["95% CI for mean", fmt_ci(H["mean_ci_electronics"]), fmt_ci(H["mean_ci_fashion"])],
              ["Skewness", f"{H['skew_electronics']:.2f}", f"{H['skew_fashion']:.2f}"]], [5, 6, 6]),
         P("3. Choice of test", H1),
         P(f"Two independent groups and a numeric outcome call for a two-sample t-test. Group sizes differ and variances are not clearly equal (Levene p = {H['levene_p']:.3f}), so the <b>Welch t-test</b> is used. "
           "Order values are right-skewed, so a Mann–Whitney test and a Welch test on log-transformed values are run as robustness checks; with n above 150 per group the t-test remains reasonable by the central limit theorem."),
         P("4. Results", H1),
         tbl([["Test", "Statistic", "p-value", "Reading"],
              ["Welch t-test (primary, one-sided)", f"t = {H['welch_t']:.2f}, df = {H['welch_df']:.0f}", f"{H['welch_p']:.3f}", "Not below 0.05"],
              ["Mann–Whitney U (one-sided)", "rank-based", f"{H['mannwhitney_p']:.3f}", "Not significant"],
              ["Welch on log(sales) (one-sided)", "log scale", f"{H['log_welch_p']:.3f}", "Not significant"],
              ["ANOVA, all 5 categories", f"F = {H['anova_F']:.2f}", f"{H['anova_p']:.3f}", "No category differs"],
              ["Kruskal–Wallis, all 5 categories", "rank-based", f"{H['kruskal_p']:.3f}", "No category differs"]], [6, 4, 2.5, 4.5]),
         Spacer(1, 6),
         tbl([["Effect size and intervals", "Value"], ["Mean difference (Electronics − Fashion)", f"{Kf(H['diff'])} ({pct(H['diff_pct'], 1)} of Fashion mean)"],
              ["95% CI for difference (Welch)", fmt_ci(H["diff_ci95"])], ["95% CI for difference (bootstrap, 10,000 resamples)", fmt_ci(H["bootstrap_ci95"])],
              ["One-sided 95% lower bound", Kf(H["diff_lower_one_sided"])], ["Cohen's d", f"{H['cohens_d']:.2f} (small)"],
              ["Bootstrap share of resamples with Electronics higher", pct(H["bootstrap_prob_gt0"], 1)], ["Orders per group for 80% power at this effect size", H["n_per_group_for_80pct_power"]]], [9, 8]),
         Spacer(1, 6), img("hypothesis_test.png", 16),
         P("5. Decision and interpretation", H1),
         P(f"<b>{H['decision']}.</b> With p = {H['welch_p']:.3f} (above α = 0.05), the data do not provide strong enough evidence that Electronics orders are larger than Fashion orders. "
           f"The observed gap of {Kf(H['diff'])} is plausible under random variation: the 95% confidence interval runs from {Kf(H['diff_ci95'][0])} to {Kf(H['diff_ci95'][1])} and includes zero. "
           "The result would be called significant at a 10% level, but that is a looser standard and was not our pre-set threshold. Failing to reject H0 is not proof that the groups are equal; the test may simply lack power for an effect this small."),
         P("<b>Business conclusion:</b> Electronics' revenue lead comes mainly from selling more orders (two products, 354 orders) rather than larger orders. Do not base pricing or inventory decisions on an assumed order-size premium without more data."),
         P("6. Limitations", H1),
         P(f"(a) About {H['n_per_group_for_80pct_power']} orders per category are needed to detect an effect this size reliably. (b) The data are observational, from one year. "
           "(c) Average unit price is nearly identical across all six products (about ₹24.7K to ₹26.4K, including Rice and Laptop), which is unusual for real retail; the data source and price field should be confirmed. "
           "(d) Dataset currency is assumed to be INR.")]
build(ROOT / "report" / "Hypothesis_Test_Report.pdf", story)

# ========================= Business_Insights.docx =========================
d = Document()
d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)
d.add_heading("Business Insights and Recommendations", 0)
d.add_paragraph(f"ApexPlanet 2025 sales: {M(K['Total Revenue'])} revenue, {K['Total Orders']:,} orders, {K['Unique Customers']} customers. Impact figures are illustrative scenarios, not forecasts.")
d.add_heading("Insight to impact", 1)
rows = [("Insight", "Problem", "Possible solution", "Recommendation", "Expected impact"),
        (f"Aug–Sep revenue ({M(aug)}, {M(sep)}) is well below the {M(avg_m)} monthly average", "Two weak months", "Seasonal offers, bundles", "Run an Aug–Sep campaign", f"About {M(dip_gain)} if both months reach average"),
        (f"{pct(1 - K['Repeat Customer Rate'])} of customers buy once; top 20% give {pct(seg['High Value']['Revenue_Share'])} of revenue", "Weak retention, revenue concentration", "Loyalty tiers, reorder reminders, VIP service", "Launch a retention programme", f"Repeat rate 5% to 10% adds about {M(retention_gain)}"),
        (f"Electronics gives {pct(K['Top Category Revenue Share'])} of revenue through volume; order-size gap not significant (p = {H['welch_p']:.3f})", "Risk of acting on an unproven price premium", "Keep stock and marketing strong; gather more data", "Back Electronics by volume", f"Avoids unproven pricing; re-test with about {H['n_per_group_for_80pct_power']} orders per category"),
        (f"Pune has high AOV ({Kf(city['Pune']['AOV'])}) but only {city['Pune']['Orders']} orders", "Under-penetrated market", "Local acquisition campaigns", "Grow Pune", f"About {M(pune_gain)} at top-3 city volume")]
t = d.add_table(rows=len(rows), cols=5); t.style = "Light Grid Accent 1"
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        t.cell(i, j).text = v
d.add_heading("Dashboard summary", 1)
d.add_picture(str(ROOT / "charts" / "dashboard_summary.png"), width=Inches(6.3))
d.save(ROOT / "report" / "Business_Insights.docx")
print("reports built")
